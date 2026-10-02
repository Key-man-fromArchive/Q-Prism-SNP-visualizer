"""16:9 PowerPoint report from accepted whole-run rows.

Slides: cover, one per marker (scatter, call counts, highlighted plate map), the
whole-plate map, then the optional results table (3 blocks x 16 rows per slide).
Fonts cannot be embedded in a .pptx; every run names the Korean-capable face for
both Latin and East Asian text so a viewer with it installed renders it.
"""

import io
from collections import Counter
from dataclasses import dataclass

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app.models import AlleleLabels, AnalysisRegionContext
from app.reporting.charts import render_scatter_png
from app.reporting.result_snapshot import ResultRow, ResultSnapshot, snapshot_rows
from app.reporting.snapshot_plate import render_snapshot_plate
from app.reporting.snapshot_presentation import (
    axis_label,
    coordinate_basis,
    cycle_label,
    display_genotype,
    figure_points,
    figure_title,
    marker_scope,
    polyploid_legend,
)

FONT = "NanumGothic"
RESULT_HEADERS = ["Well", "Sample", "Marker", "Allele Call", "Confidence (%)"]
ROWS_PER_BLOCK = 16
BLOCKS_PER_SLIDE = 3
WELLS_PER_PAGE = ROWS_PER_BLOCK * BLOCKS_PER_SLIDE
_COLUMN_WIDTHS = (0.5, 1.1, 0.95, 1.0, 0.65)
_SAMPLE_CHARS, _MARKER_CHARS, _LAYOUT_CHARS = 18, 14, 10
_SLIDE_W, _SLIDE_H = 13.333, 7.5
_HEADER_FILL = RGBColor(0xE8, 0xEE, 0xF4)
_TEXT = RGBColor(0x11, 0x18, 0x27)
_BOX = tuple[float, float, float, float]


@dataclass(frozen=True)
class _MarkerPage:
    title: str
    marker: AnalysisRegionContext | None
    labels: AlleleLabels | None
    rows: list[ResultRow]


def result_page_count(well_count: int) -> int:
    """Result slides needed for ``well_count`` wells (3 blocks x 16 rows each)."""
    return -(-well_count // WELLS_PER_PAGE)


def _style_run(run, size: float, bold: bool = False) -> None:
    font = run.font
    font.size, font.bold = Pt(size), bold
    font.color.rgb = _TEXT
    font.name = FONT
    latin = run._r.rPr.find(qn("a:latin"))
    latin.addnext(latin.makeelement(qn("a:ea"), {"typeface": FONT}))


def _write(text_frame, text: object, size: float, bold: bool = False) -> None:
    paragraph = text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "" if text is None else str(text)
    _style_run(run, size, bold)


def _clip(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _textbox(slide, box: _BOX, text: str, size: float, bold: bool = False) -> None:
    left, top, width, height = box
    frame = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    ).text_frame
    frame.word_wrap = True
    _write(frame, text, size, bold)


def _fit_picture(slide, png: bytes, box: _BOX) -> None:
    left, top, width, height = box
    with Image.open(io.BytesIO(png)) as image:
        ratio = image.width / image.height
    fitted_w, fitted_h = (
        (width, width / ratio) if width / height < ratio else (height * ratio, height)
    )
    slide.shapes.add_picture(
        io.BytesIO(png),
        Inches(left + (width - fitted_w) / 2),
        Inches(top + (height - fitted_h) / 2),
        Inches(fitted_w),
        Inches(fitted_h),
    )


def _add_table(
    slide,
    header: list[str],
    body: list[list[object]],
    origin: tuple[float, float],
    widths: tuple[float, ...],
    row_height: float,
    size: float,
):
    shape = slide.shapes.add_table(
        len(body) + 1,
        len(header),
        Inches(origin[0]),
        Inches(origin[1]),
        Inches(sum(widths)),
        Inches(row_height * (len(body) + 1)),
    )
    table = shape.table
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row_index, values in enumerate([header, *body]):
        table.rows[row_index].height = Inches(row_height)
        for column, value in enumerate(values):
            cell = table.cell(row_index, column)
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            _write(cell.text_frame, value, size, bold=row_index == 0)
            if row_index == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _HEADER_FILL
    return table


def _blank(deck, title: str):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _textbox(slide, (0.4, 0.25, 12.5, 0.6), title, 20, bold=True)
    return slide


def _cover(deck, snapshot: ResultSnapshot, marker_count: int) -> None:
    context = snapshot.context
    slide = _blank(deck, "SNP Discrimination Report")
    details = [
        ["File", snapshot.raw_filename],
        ["Instrument", snapshot.unified.instrument],
        ["Analysis cycle", cycle_label(snapshot, context.cycle)],
        ["Use ROX", context.use_rox],
        ["Passive reference", snapshot.passive_reference_label],
        ["Background", context.background],
        ["Markers", marker_count],
        ["Result revision", context.result_revision],
        ["Analysed at", context.analysed_at.isoformat()],
    ]
    _add_table(slide, ["Condition", "Value"], details, (0.8, 1.2), (3.0, 8.7), 0.45, 14)


def _marker_pages(snapshot: ResultSnapshot, rows: list[ResultRow]) -> list[_MarkerPage]:
    if not snapshot.context.regions:
        return [_MarkerPage("Whole-run", None, None, rows)]
    pages = []
    for region in snapshot.context.regions:
        label = snapshot.marker_labels.get(region.marker_id)
        name = label.name if label else region.name
        own = [
            row
            for row in rows
            if row.marker and row.marker.marker_id == region.marker_id
        ]
        pages.append(
            _MarkerPage(name, region, label.allele_labels if label else None, own)
        )
    return pages


def _scatter(snapshot: ResultSnapshot, page: _MarkerPage) -> bytes:
    marker = page.marker
    points = figure_points(page.rows)
    legend_names = {
        genotype: display_genotype(genotype, marker, page.labels)
        for genotype in {point["effective_type"] for point in points}
    }
    return render_scatter_png(
        points,
        snapshot.unified.allele2_dye,
        ploidy=marker.ploidy if marker else snapshot.result.ploidy,
        coordinate_basis=coordinate_basis(snapshot),
        title=_page_title(snapshot, page),
        x_label=axis_label(snapshot, "allele2", marker, page.labels),
        y_label=axis_label(snapshot, "fam", marker, page.labels),
        legend_names=legend_names,
    )


def _page_title(snapshot: ResultSnapshot, page: _MarkerPage) -> str:
    if page.marker is None:
        return f"{page.title} · {cycle_label(snapshot, snapshot.context.cycle)}"
    return figure_title(snapshot, page.title, page.marker)


def _call_counts(page: _MarkerPage) -> list[list[object]]:
    counts = Counter(
        display_genotype(row.genotype, row.marker, row.allele_labels)
        for row in page.rows
    )
    return [[call, count] for call, count in counts.items()]


def _plate_map(
    rows: list[ResultRow], layout: dict[str, str], highlight: str | None = None
) -> bytes | None:
    if not rows:
        return None
    return render_snapshot_plate(
        rows, marker_layout=layout or None, highlight_marker_id=highlight
    )


def _marker_slide(
    deck,
    snapshot: ResultSnapshot,
    page: _MarkerPage,
    rows: list[ResultRow],
    layout: dict[str, str],
) -> None:
    marker = page.marker
    slide = _blank(deck, _page_title(snapshot, page))
    legend = (
        polyploid_legend(page.labels, snapshot.unified.allele2_dye)
        if marker and marker.ploidy != 2
        else None
    )
    if legend:
        _textbox(slide, (0.4, 0.8, 12.5, 0.3), legend, 11)
    _fit_picture(slide, _scatter(snapshot, page), (0.4, 1.15, 7.5, 6.05))
    _add_table(
        slide,
        ["Allele Call", "Count"],
        _call_counts(page)[:10],
        (8.2, 1.15),
        (3.4, 1.3),
        0.26,
        10,
    )
    plate = _plate_map(rows, layout, marker.marker_id if marker else None)
    if plate:
        _fit_picture(slide, plate, (8.2, 4.2, 4.7, 3.0))


def _result_table_rows(rows: list[ResultRow]) -> list[list[object]]:
    return [
        [
            row.well,
            _clip(row.sample_name, _SAMPLE_CHARS),
            _clip(row.marker.name if row.marker else "", _MARKER_CHARS),
            display_genotype(row.genotype, row.marker, row.allele_labels),
            round(row.confidence * 100, 1) if row.confidence is not None else "",
        ]
        for row in rows
    ]


def _result_slides(deck, rows: list[ResultRow]) -> None:
    body = _result_table_rows(rows)
    pages = result_page_count(len(body))
    for page in range(pages):
        slide = _blank(deck, f"Results ({page + 1}/{pages})")
        chunk = body[page * WELLS_PER_PAGE : (page + 1) * WELLS_PER_PAGE]
        for block in range(BLOCKS_PER_SLIDE):
            part = chunk[block * ROWS_PER_BLOCK : (block + 1) * ROWS_PER_BLOCK]
            if part:
                _add_table(
                    slide,
                    RESULT_HEADERS,
                    part,
                    (0.35 + block * 4.3, 0.95),
                    _COLUMN_WIDTHS,
                    0.3,
                    8,
                )


def build_snapshot_pptx(snapshot: ResultSnapshot, include_table: bool = True) -> bytes:
    rows = snapshot_rows(snapshot)
    pages = _marker_pages(snapshot, rows)
    layout = {
        page.marker.marker_id: _clip(page.title, _LAYOUT_CHARS)
        for page in pages
        if page.marker
    }
    deck = Presentation()
    deck.slide_width, deck.slide_height = (
        Emu(int(_SLIDE_W * 914400)),
        Emu(int(_SLIDE_H * 914400)),
    )
    _cover(deck, snapshot, len(pages))
    for page in pages:
        _marker_slide(deck, snapshot, page, rows, layout)
    plate_slide = _blank(deck, "Plate view")
    plate = _plate_map(rows, layout)
    if plate:
        _fit_picture(plate_slide, plate, (0.4, 1.0, 12.5, 6.2))
    if include_table:
        _result_slides(deck, rows)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def pptx_filename(snapshot: ResultSnapshot, selected: bool) -> str:
    """Download name; a marker selection puts the (safely reduced) marker names in it."""
    scope = marker_scope(snapshot, selected) or "whole-run"
    return f"snp_report_{scope}_cycle{snapshot.context.cycle}.pptx"
