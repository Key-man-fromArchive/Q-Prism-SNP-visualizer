"""Typed XLSX cells and real figures from one accepted result."""
import io
from fastapi import HTTPException

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.reporting.charts import render_scatter_png
from app.reporting.result_snapshot import ResultSnapshot, snapshot_rows
from app.reporting.snapshot_presentation import (
    CellValue, coordinate_basis, report_counts, report_figures, report_metadata, report_table,
)


def _validate_text(values: list[CellValue]) -> None:
    if any(isinstance(value, str) and len(value) > 32767 for value in values):
        raise HTTPException(400, "XLSX text exceeds 32767 characters; export CSV for full labels")


def append_cells(sheet: Worksheet, values: list[CellValue]) -> None:
    """Prevent formula inference without modifying labels or negative numbers."""
    _validate_text(values)
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_snapshot_xlsx(snapshot: ResultSnapshot) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    if summary is None:
        raise RuntimeError("Workbook has no summary sheet")
    summary.title = "Summary"
    append_cells(summary, ["SNP Allele Discrimination Report"])
    for key, value in report_metadata(snapshot):
        append_cells(summary, [key, value])
    append_cells(summary, ["Marker ID", "Genotype", "Count"])
    rows = snapshot_rows(snapshot)
    _summary_qc(summary, [row.genotype for row in rows])
    for count in report_counts(rows):
        append_cells(summary, count)
    for index, figure in enumerate(report_figures(snapshot, rows)):
        anchor = 3 + index * 35
        _validate_text([figure.title])
        summary.cell(anchor, 4, figure.title).data_type = "s"
        png = render_scatter_png(figure.points, snapshot.unified.allele2_dye, ploidy=figure.ploidy,
                                 coordinate_basis=coordinate_basis(snapshot))
        image = Image(io.BytesIO(png))
        image.width, image.height = 640, 480
        summary.add_image(image, f"D{anchor + 1}")
    results = workbook.create_sheet("Results")
    headers, values = report_table(snapshot)
    append_cells(results, list(headers))
    for row_values in values:
        append_cells(results, row_values)
    for cell in results[1]:
        cell.font = Font(bold=True)
    results.freeze_panes = "A2"
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 55
    for column in ("A", "B", "C", "D"):
        results.column_dimensions[column].width = 24
    context = workbook.create_sheet("Analysis Context")
    append_cells(context, ["Chunk", "Analysis Context (concatenate in chunk order)"])
    encoded = snapshot.context.model_dump_json()
    for index, start in enumerate(range(0, len(encoded), 30000), 1):
        append_cells(context, [index, encoded[start:start + 30000]])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _summary_qc(summary: Worksheet, genotypes: list[str]) -> None:
    # Preserve the pre-contract report's counting policy for existing types.
    called = sum(kind not in ("Unknown", "Unassigned", "Undetermined", "NTC") for kind in genotypes)
    append_cells(summary, ["Total wells", len(genotypes)])
    append_cells(summary, ["Called", called])
    append_cells(summary, ["Call rate (%)", round(100 * called / len(genotypes), 1) if genotypes else 0])
    append_cells(summary, ["Called policy", "Legacy report: excludes Unknown, Undetermined, NTC; also excludes new Unassigned. Empty/Omit retain legacy inclusion."])
