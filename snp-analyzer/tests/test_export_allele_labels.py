"""P2-E2: allele names in CSV, XLSX and QC (synthetic StepOnePlus upload -> output)."""

# ruff: noqa: F811  (the imported restore fixture is re-declared as a test argument)
import csv
import io

from openpyxl import load_workbook

from tests import stepone_fixtures as so
from tests.test_stepone_session_restore import restore  # noqa: F401

CLUSTER = {"cycle": 2, "use_rox": True, "algorithm": "threshold"}
CANONICAL = {"Allele 1 Homo", "Heterozygous", "Allele 2 Homo"}
DISPLAY = {"WT/WT", "WT/MT", "MT/MT"}


def _ready(client, preset: str = "partial_default") -> str:
    options = so.StepOneOptions(markers=so.default_markers(preset))
    response = client.post(
        "/api/upload",
        files={"file": ("plate.eds", so.build_stepone_eds(options), "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    sid = response.json()["session_id"]
    assert client.post(f"/api/data/{sid}/cluster", json=CLUSTER).status_code == 200
    return sid


def _marker_id(client, sid: str, name: str) -> str:
    markers = client.get(f"/api/data/{sid}/markers").json()["markers"]
    return next(m["id"] for m in markers if m["name"] == name)


def _csv_rows(client, sid: str) -> tuple[list[str], list[dict[str, str]], object]:
    response = client.get(f"/api/data/{sid}/export/csv")
    assert response.status_code == 200, response.text
    reader = csv.DictReader(io.StringIO(response.text))
    return list(reader.fieldnames or []), list(reader), response


def _xlsx(client, sid: str):
    response = client.get(f"/api/data/{sid}/export/xlsx")
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content)), response


def test_named_csv_has_allele_call_after_canonical_genotype(restore):
    client, _, _ = restore
    sid = _ready(client)
    header, rows, _ = _csv_rows(client, sid)
    assert header[header.index("Genotype") + 1] == "Allele Call"
    named = [r for r in rows if r["Marker"] == "QPrism1" and r["Genotype"] in CANONICAL]
    assert named
    assert {r["Allele Call"] for r in named} <= DISPLAY
    assert all(r["Genotype"] in CANONICAL for r in named)
    plain = [r for r in rows if r["Marker"] == "QPrism2"]
    assert all(r["Allele Call"] == r["Genotype"] for r in plain)


def test_csv_has_cycle_label_for_a_stepone_run(restore):
    client, _, _ = restore
    sid = _ready(client)
    header, rows, _ = _csv_rows(client, sid)
    assert header[header.index("Cycle") + 1] == "Cycle Label"
    assert "PCR" in rows[0]["Cycle Label"]


def test_unnamed_csv_has_no_allele_call_column(restore):
    client, _, _ = restore
    sid = _ready(client, "all_default")
    header, _, _ = _csv_rows(client, sid)
    assert "Allele Call" not in header


def test_csv_allele_call_is_formula_safe(restore):
    client, _, _ = restore
    sid = _ready(client)
    marker_id = _marker_id(client, sid, "QPrism2")
    response = client.put(
        f"/api/data/{sid}/markers/{marker_id}",
        json={"allele_labels": {"fam": "=WT", "allele2": "+MT"}},
    )
    assert response.status_code == 200, response.text
    assert client.post(f"/api/data/{sid}/cluster", json=CLUSTER).status_code == 200
    _, rows, _ = _csv_rows(client, sid)
    calls = {r["Allele Call"] for r in rows if r["Marker"] == "QPrism2"}
    dangerous = {c for c in calls if "=WT" in c or "+MT" in c}
    assert dangerous
    assert all(c.startswith("'") for c in dangerous)


def test_downloads_use_rfc5987_headers(restore):
    client, _, _ = restore
    sid = _ready(client)
    csv_header = client.get(f"/api/data/{sid}/export/csv").headers["content-disposition"]
    xlsx_header = client.get(f"/api/data/{sid}/export/xlsx").headers["content-disposition"]
    assert "filename*=UTF-8''snp_export_whole-run_cycle2.csv" in csv_header
    assert "filename*=UTF-8''snp_report_whole-run_cycle2.xlsx" in xlsx_header


def test_xlsx_results_allele_call_and_cycle_label(restore):
    client, _, _ = restore
    sid = _ready(client)
    workbook, _ = _xlsx(client, sid)
    rows = list(workbook["Results"].iter_rows(values_only=True))
    header = list(rows[0])
    assert header[header.index("Genotype") + 1] == "Allele Call"
    assert header[header.index("Cycle") + 1] == "Cycle Label"
    genotype, call = header.index("Genotype"), header.index("Allele Call")
    marker = header.index("Marker")
    named = [r for r in rows[1:] if r[marker] == "QPrism1" and r[genotype] in CANONICAL]
    assert named and {r[call] for r in named} <= DISPLAY


def test_xlsx_summary_has_one_figure_per_marker_with_allele_axes(restore):
    client, _, _ = restore
    sid = _ready(client)
    workbook, _ = _xlsx(client, sid)
    summary = workbook["Summary"]
    assert len(summary._images) == 6
    cells = [c for row in summary.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert any("QPrism1" in c for c in cells)
    assert any(c.startswith("Analysis read") or "PCR" in c for c in cells)


def test_unnamed_xlsx_has_no_allele_call_column(restore):
    client, _, _ = restore
    sid = _ready(client, "all_default")
    workbook, _ = _xlsx(client, sid)
    header = [c.value for c in workbook["Results"][1]]
    assert "Allele Call" not in header


def test_xlsx_allele_call_cells_stay_text(restore):
    client, _, _ = restore
    sid = _ready(client)
    marker_id = _marker_id(client, sid, "QPrism2")
    client.put(
        f"/api/data/{sid}/markers/{marker_id}",
        json={"allele_labels": {"fam": "=WT", "allele2": "+MT"}},
    )
    assert client.post(f"/api/data/{sid}/cluster", json=CLUSTER).status_code == 200
    workbook, _ = _xlsx(client, sid)
    results = workbook["Results"]
    header = [c.value for c in results[1]]
    column = header.index("Allele Call")
    cells = [row[column] for row in results.iter_rows(min_row=2)]
    formulas = [c for c in cells if isinstance(c.value, str) and c.value.startswith(("=", "+"))]
    assert formulas
    assert all(c.data_type == "s" for c in formulas)


def test_qc_markers_carry_display_names_and_allele_names(restore):
    client, _, _ = restore
    sid = _ready(client)
    qc = client.get(f"/api/data/{sid}/qc?cycle=2").json()
    by_name = {m["name"]: m for m in qc["markers"]}
    assert by_name["QPrism1"]["allele_labels"] == {"fam": "WT", "allele2": "MT"}
    assert by_name["QPrism1"]["display_name"] == "QPrism1"
    assert "allele_labels" not in by_name["QPrism2"]


def test_qc_display_name_follows_a_rename(restore):
    client, _, _ = restore
    sid = _ready(client)
    marker_id = _marker_id(client, sid, "QPrism3")
    client.put(f"/api/data/{sid}/markers/{marker_id}", json={"name": "한글마커"})
    qc = client.get(f"/api/data/{sid}/qc?cycle=2").json()
    renamed = next(m for m in qc["markers"] if m["id"] == marker_id)
    assert renamed["display_name"] == "한글마커"


def test_unnamed_qc_markers_are_unchanged(restore):
    client, _, _ = restore
    sid = _ready(client, "all_default")
    qc = client.get(f"/api/data/{sid}/qc?cycle=2").json()
    for marker in qc["markers"]:
        assert "allele_labels" not in marker and "display_name" not in marker
