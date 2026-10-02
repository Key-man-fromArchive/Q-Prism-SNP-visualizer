"""PPTX report: StepOnePlus upload -> analysis -> /export/pptx, inspected with python-pptx."""

import io
import math
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Emu

from stepone_fixtures import build_stepone_eds
from test_marker_contract import data_client as _data_client

data_client = _data_client

URL = "/api/data/{sid}/export/pptx"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
CANONICAL_DIPLOID = {"Allele 1 Homo", "Heterozygous", "Allele 2 Homo"}


@pytest.fixture
def run(data_client: SimpleNamespace) -> SimpleNamespace:
    """A synthetic StepOnePlus run (6 markers x 16 wells, QPrism1 named) analysed once."""
    client = data_client.client
    conn = data_client.db.get_db()
    conn.execute(
        "INSERT OR IGNORE INTO users (id, username, hashed_password, display_name, role) "
        "VALUES ('user-1', 'pptx-user', 'x', 'pptx-user', 'user')"
    )
    conn.commit()
    response = client.post(
        "/api/upload",
        files={"file": ("plate.eds", build_stepone_eds(), "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    sid = body["session_id"]
    markers = client.get(f"/api/data/{sid}/markers").json()["markers"]
    assert len(markers) == 6
    analysed = client.post(
        f"/api/data/{sid}/cluster",
        json={
            "cycle": body["suggested_cycle"],
            "use_rox": False,
            "algorithm": "threshold",
        },
    )
    assert analysed.status_code == 200, analysed.text
    return SimpleNamespace(client=client, sid=sid, markers=markers, base=data_client)


def _deck(response) -> Presentation:
    assert response.status_code == 200, response.text
    return Presentation(io.BytesIO(response.content))


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(parts)


def _tables(slide):
    return [shape.table for shape in slide.shapes if shape.has_table]


def _pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == 13]


def _result_slides(deck, marker_count: int):
    return list(deck.slides)[2 + marker_count :]


def test_slide_count_is_cover_markers_plate_and_result_pages(run) -> None:
    response = run.client.get(URL.format(sid=run.sid))
    assert response.headers["content-type"].startswith(PPTX_TYPE)
    deck = _deck(response)
    # 96 wells -> 2 result slides of 3 x 16 = 48 wells.
    assert len(deck.slides) == 1 + 6 + 1 + 2


def test_deck_is_widescreen(run) -> None:
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    assert math.isclose(deck.slide_width / deck.slide_height, 16 / 9, rel_tol=0.01)


def test_results_table_can_be_excluded(run) -> None:
    deck = _deck(
        run.client.get(URL.format(sid=run.sid), params={"include_table": "false"})
    )
    assert len(deck.slides) == 1 + 6 + 1
    assert not any(
        _tables(slide) and "Allele Call" in _slide_text(slide)
        for slide in list(deck.slides)[8:]
    )


def test_marker_selection_limits_marker_and_result_slides(run) -> None:
    ids = ",".join(m["id"] for m in run.markers[:2])
    deck = _deck(run.client.get(URL.format(sid=run.sid), params={"marker_ids": ids}))
    # 32 wells fit one result slide.
    assert len(deck.slides) == 1 + 2 + 1 + 1
    text = "\n".join(_slide_text(s) for s in _result_slides(deck, 2))
    assert run.markers[0]["name"] in text and run.markers[5]["name"] not in text


def test_marker_slide_has_scatter_map_and_count_table(run) -> None:
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    slide = list(deck.slides)[1]
    assert run.markers[0]["name"] in _slide_text(slide)
    assert len(_pictures(slide)) == 2
    counts = _tables(slide)
    assert counts and counts[0].rows[0].cells[0].text


def test_cover_slide_carries_provenance(run) -> None:
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    text = _slide_text(list(deck.slides)[0])
    revision = run.base.clustering.cluster_store[
        run.sid
    ].analysis_context.result_revision
    assert "plate.eds" in text and str(revision) in text
    assert "StepOne" in text


def test_plate_slide_has_the_map(run) -> None:
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    assert len(_pictures(list(deck.slides)[7])) == 1


def test_result_pages_hold_three_blocks_of_sixteen_rows(run) -> None:
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    slides = _result_slides(deck, 6)
    for slide in slides:
        blocks = _tables(slide)
        assert len(blocks) == 3
        for block in blocks:
            assert [c.text for c in block.rows[0].cells] == [
                "Well",
                "Sample",
                "Marker",
                "Allele Call",
                "Confidence (%)",
            ]
            assert len(block.rows) == 1 + 16
    wells = [
        row.cells[0].text
        for slide in slides
        for block in _tables(slide)
        for row in list(block.rows)[1:]
        if row.cells[0].text
    ]
    assert len(wells) == 96 and len(set(wells)) == 96


def _result_rows(deck, marker_count=6):
    return [
        [c.text for c in row.cells]
        for slide in _result_slides(deck, marker_count)
        for block in _tables(slide)
        for row in list(block.rows)[1:]
        if row.cells[0].text
    ]


def test_allele_call_uses_names_only_for_named_markers(run) -> None:
    rows = _result_rows(_deck(run.client.get(URL.format(sid=run.sid))))
    first = [r for r in rows if r[2] == "QPrism1"]
    other = [r for r in rows if r[2] != "QPrism1"]
    assert len(first) == 16 and len(other) == 80
    named = {"WT/WT", "WT/MT", "MT/MT"}
    assert {r[3] for r in first} & named
    assert not {r[3] for r in first} & CANONICAL_DIPLOID
    assert not {r[3] for r in other} & named
    assert {r[3] for r in other} & CANONICAL_DIPLOID


def test_rename_after_analysis_is_reflected(run) -> None:
    marker = run.markers[0]
    response = run.client.put(
        f"/api/data/{run.sid}/markers/{marker['id']}",
        json={"name": "Renamed", "allele_labels": {"fam": "REF", "allele2": "ALT"}},
    )
    assert response.status_code == 200, response.text
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    assert "Renamed" in _slide_text(list(deck.slides)[1])
    rows = _result_rows(deck)
    assert {r[3] for r in rows if r[2] == "Renamed"} & {"REF/REF", "REF/ALT", "ALT/ALT"}
    assert not [r for r in rows if r[2] == "QPrism1"]


def test_every_run_names_a_korean_font_for_latin_and_east_asian(run) -> None:
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    runs = 0
    for slide in deck.slides:
        for element in slide._element.iter("{%s}rPr" % NS["a"]):
            latin = element.find("a:latin", NS)
            east_asian = element.find("a:ea", NS)
            assert latin is not None and east_asian is not None
            assert latin.get("typeface") == east_asian.get("typeface") != ""
            runs += 1
    assert runs > 100


def test_text_stays_inside_the_slide(run) -> None:
    deck = _deck(run.client.get(URL.format(sid=run.sid)))
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width + Emu(1)
            assert shape.top + shape.height <= deck.slide_height + Emu(1)


def test_download_header_is_rfc5987_and_names_the_run(run) -> None:
    header = run.client.get(URL.format(sid=run.sid)).headers["content-disposition"]
    assert header.startswith("attachment; filename=")
    assert ".pptx" in header and "filename*=UTF-8''" in header


def test_selected_marker_name_in_filename_is_safe(run) -> None:
    run.client.put(
        f"/api/data/{run.sid}/markers/{run.markers[0]['id']}",
        json={"name": "../évil\x01:name"},
    )
    header = run.client.get(
        URL.format(sid=run.sid), params={"marker_ids": run.markers[0]["id"]}
    ).headers["content-disposition"]
    assert "/" not in header.split("filename=")[1].split(";")[0]
    assert ".." not in header.split("filename=")[1].split(";")[0]
    assert "\x01" not in header


@pytest.mark.parametrize(
    "query", ["cycle=3&cycle_mode=absolute", "use_rox=true", "background=pre_read"]
)
def test_condition_mismatch_is_409(run, query: str) -> None:
    response = run.client.get(f"{URL.format(sid=run.sid)}?{query}")
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "EXPORT_CONDITION_MISMATCH"


def test_stale_result_revision_is_409(run) -> None:
    response = run.client.get(
        URL.format(sid=run.sid),
        params={"result_revision": "00000000-0000-0000-0000-000000000000"},
    )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "RESULT_REVISION_CONFLICT"


@pytest.mark.parametrize("marker_ids", ["", "nope", "a,,b", "x,x"])
def test_bad_marker_ids_are_400(run, marker_ids: str) -> None:
    response = run.client.get(
        URL.format(sid=run.sid), params={"marker_ids": marker_ids}
    )
    assert response.status_code == 400


def test_marker_ids_are_checked_before_the_result_state(run) -> None:
    run.base.clustering.cluster_store.pop(run.sid)
    assert (
        run.client.get(URL.format(sid=run.sid), params={"marker_ids": ""}).status_code
        == 400
    )
    assert run.client.get(URL.format(sid=run.sid)).status_code == 409


def test_unknown_session_is_404(run) -> None:
    assert run.client.get(URL.format(sid="missing")).status_code == 404


def test_unnamed_run_matches_canonical_calls(run) -> None:
    """Markers without names keep the canonical call string, never a blank."""
    rows = _result_rows(_deck(run.client.get(URL.format(sid=run.sid))))
    assert all(r[3] for r in rows)


def test_result_page_count_helper() -> None:
    from app.reporting.snapshot_pptx import result_page_count

    assert [result_page_count(n) for n in (0, 1, 48, 49, 96, 384)] == [0, 1, 1, 2, 2, 8]
