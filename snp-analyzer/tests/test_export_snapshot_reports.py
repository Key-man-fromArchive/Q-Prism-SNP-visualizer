"""Report formats obey the same accepted result as the CSV adapter."""
import csv
import io
import os
from uuid import UUID
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from test_export_snapshot_csv import plate, data_client  # noqa: F401


@pytest.fixture(autouse=True)
def isolated_report_metadata(monkeypatch: pytest.MonkeyPatch, data_client: SimpleNamespace) -> None:
    """Replace only metadata dictionaries; monkeypatch restores prior state."""
    from app.routers import sample, data
    from app import asg_session
    monkeypatch.setattr(sample, "sample_name_store", {})
    monkeypatch.setattr(data, "protocol_store", {})
    monkeypatch.setattr(asg_session, "_current_launch_by_user", {})
    monkeypatch.setattr(asg_session, "_launch_by_session", {})


@pytest.mark.parametrize("extension", ["pdf", "xlsx"])
@pytest.mark.parametrize("state,code", [
    ("missing", "NO_COMPLETED_RESULT"),
    ("legacy", "LEGACY_CONTEXT_UNKNOWN"),
    ("stale", "INPUT_REVISION_CONFLICT"),
    ("mismatch", "EXPORT_CONDITION_MISMATCH"),
])
def test_report_rejects_unverified_conditions(plate: SimpleNamespace, extension: str, state: str, code: str) -> None:
    result = plate.clustering.cluster_store["export"]
    query = ""
    if state == "missing":
        plate.clustering.cluster_store.pop("export")
    elif state == "legacy":
        result.analysis_context = None
    elif state == "stale":
        plate.upload.sessions["export"].input_revision += 1
    else:
        query = "?cycle=40"
    response = plate.client.get(f"/api/data/export/export/{extension}{query}")
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == code


@pytest.mark.parametrize("extension", ["csv", "pdf", "xlsx"])
def test_explicit_absolute_zero_is_not_legacy_latest(plate: SimpleNamespace, extension: str) -> None:
    response = plate.client.get(f"/api/data/export/export/{extension}?cycle=0&cycle_mode=absolute")
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "EXPORT_CONDITION_MISMATCH"


@pytest.mark.parametrize("extension", ["csv", "pdf", "xlsx"])
def test_sparse_actual_zero_is_exported_with_absolute_mode(plate: SimpleNamespace, extension: str) -> None:
    clustered = plate.client.post("/api/data/export/cluster", json={
        "cycle": 0, "cycle_mode": "absolute", "use_rox": False, "algorithm": "threshold",
    })
    assert clustered.status_code == 200, clustered.text
    revision = clustered.json()["analysis_context"]["result_revision"]
    response = plate.client.get(
        f"/api/data/export/export/{extension}?cycle=0&cycle_mode=absolute&result_revision={revision}"
    )
    assert response.status_code == 200, response.text
    assert "cycle0" in response.headers["content-disposition"]


def test_asg_snapshot_uses_sparse_actual_zero_with_absolute_mode(plate: SimpleNamespace) -> None:
    from app.asg_result import build_result_snapshot
    from app.auth import TokenData

    clustered = plate.client.post("/api/data/export/cluster", json={
        "cycle": 0, "cycle_mode": "absolute", "use_rox": False, "algorithm": "threshold",
    })
    assert clustered.status_code == 200, clustered.text
    revision = clustered.json()["analysis_context"]["result_revision"]
    _bind_report_asg()

    payload = build_result_snapshot(
        "export", user=TokenData(user_id="u", username="u", role="user"), selected_cycle=0,
        cycle_mode="absolute", result_revision=UUID(revision),
    )

    assert payload["selected_cycle"] == 0
    assert payload["result"]["analysis_context"]["cycle"] == 0
    assert payload["result"]["analysis_context"]["result_revision"] == revision


def test_xlsx_values_match_csv_and_contain_plot(plate: SimpleNamespace) -> None:
    csv_response = plate.client.get("/api/data/export/export/csv")
    expected = list(csv.DictReader(io.StringIO(csv_response.text)))
    response = plate.client.get("/api/data/export/export/xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Results"]
    headers = [cell.value for cell in sheet[1]]
    actual = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert len(actual) == len(expected) == 96
    for report, source in zip(actual, expected):
        for key in ("Well", "Genotype", "Result Revision", "Read Status", "Assignment Status"):
            assert report[key] == source[key]
        assert report["FAM (norm)"] == float(source["FAM (norm)"])
        assert report["Cycle"] == 20
    assert workbook["Summary"]._images
    assert "whole-run_cycle20" in response.headers["content-disposition"]


def test_pdf_contains_stored_values_and_provenance(plate: SimpleNamespace) -> None:
    response = plate.client.get("/api/data/export/export/pdf")
    assert response.status_code == 200
    pdf = PdfReader(io.BytesIO(response.content))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    revision = str(plate.clustering.cluster_store["export"].analysis_context.result_revision)
    for value in (revision, "whole-run", "Analysis cycle", "20", "90.0", "Confidence", "A1", "H12", "Plate view", "Full-curve Ct", "background=none"):
        assert value in text
    assert "whole-run_cycle20" in response.headers["content-disposition"]


def test_asg_uses_same_snapshot_and_rejects_legacy(plate: SimpleNamespace) -> None:
    from app.asg_client import ASGLaunchContext, ASGLaunchSaveCredential
    from app.asg_session import bind_session_to_current_asg_launch, remember_asg_launch
    from app.asg_result import build_result_snapshot
    from app.auth import TokenData
    from fastapi import HTTPException

    remember_asg_launch("u", ASGLaunchContext("design_run_item", "1", {}),
                        ASGLaunchSaveCredential("launch", "test-token"), ["snp:save_result"], None)
    bind_session_to_current_asg_launch("export", "u")
    user = TokenData(user_id="u", username="u", role="user")
    payload = build_result_snapshot("export", user=user)
    assert payload["selected_cycle"] == 20
    assert payload["result"]["wells"][0]["norm_fam"] == 90.0
    assert payload["result"]["analysis_context"]["cycle"] == 20
    assert "data" not in payload["result"]
    plate.clustering.cluster_store["export"].analysis_context = None
    with pytest.raises(HTTPException) as error:
        build_result_snapshot("export", user=user)
    assert error.value.detail["code"] == "LEGACY_CONTEXT_UNKNOWN"


def test_xlsx_preserves_large_context_and_formula_text(plate: SimpleNamespace) -> None:
    from app.routers.sample import sample_name_store
    result = plate.clustering.cluster_store["export"]
    result.analysis_context.parameters["long_note"] = "가" * 40000
    sample_name_store["export"] = {"A1": "=SUM(1,2)"}
    response = plate.client.get("/api/data/export/export/xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    chunks = list(workbook["Analysis Context"].iter_rows(min_row=2, values_only=True))
    assert "".join(row[1] for row in chunks) == result.analysis_context.model_dump_json()
    names = [cell.value for cell in workbook["Results"][1]]
    cell = workbook["Results"].cell(2, names.index("Sample Name") + 1)
    assert cell.value == "=SUM(1,2)" and cell.data_type == "s"
    summary = list(workbook["Summary"].values)
    assert any(row[0] == "Total wells" and row[1] == 96 for row in summary)


def test_xlsx_rejects_oversized_label_instead_of_truncating(plate: SimpleNamespace) -> None:
    from app.routers.sample import sample_name_store
    sample_name_store["export"] = {"A1": "x" * 33000}
    response = plate.client.get("/api/data/export/export/xlsx")
    assert response.status_code == 400
    assert "32767" in response.json()["detail"]


def test_korean_report_embeds_font_and_renders_all_pages(plate: SimpleNamespace, tmp_path: Path) -> None:
    import pypdfium2 as pdfium
    from app.routers.sample import sample_name_store
    label = "긴한글시료이름" * 12 + "<안전>&끝"
    sample_name_store["export"] = {"A1": label}
    response = plate.client.get("/api/data/export/export/pdf")
    assert response.status_code == 200
    document = PdfReader(io.BytesIO(response.content))
    text = "".join(page.extract_text() for page in document.pages).replace("\n", "").replace(" ", "")
    assert label in text
    fonts = [font.get_object() for page in document.pages
             for font in page["/Resources"]["/Font"].get_object().values()]
    assert any("/FontFile2" in font.get("/FontDescriptor", {}).get_object()
               for font in fonts if "/FontDescriptor" in font)
    directory = Path(os.environ.get("UX_REPORT_ARTIFACTS", str(tmp_path)))
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "korean-report.pdf").write_bytes(response.content)
    with pdfium.PdfDocument(response.content) as pdf:
        assert len(pdf) >= 5
        for index in range(len(pdf)):
            page = pdf[index]
            bitmap = page.render(scale=1.5)
            image = bitmap.to_pil()
            assert image.width > 1000 and image.height > 700
            image.save(directory / f"korean-page-{index + 1}.png")
            bitmap.close()
            page.close()


@pytest.mark.parametrize("extension", ["pdf", "xlsx"])
def test_report_figure_uses_captured_coordinates_and_basis(plate: SimpleNamespace, monkeypatch: pytest.MonkeyPatch, extension: str) -> None:
    from app.reporting import snapshot_pdf, snapshot_xlsx
    module = snapshot_pdf if extension == "pdf" else snapshot_xlsx
    original = module.render_scatter_png
    captured = []
    def render(points, dye, **kwargs):
        captured.append((points, kwargs))
        return original(points, dye, **kwargs)
    monkeypatch.setattr(module, "render_scatter_png", render)
    response = plate.client.get(f"/api/data/export/export/{extension}")
    assert response.status_code == 200
    points, settings = captured[0]
    assert points[0]["norm_fam"] == 90.0
    assert settings["coordinate_basis"] == "raw / post-background"
    if extension == "pdf":
        text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
        assert "ClusteringAlgorithm." not in text


def test_plate_384_has_complete_extent(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.reporting import snapshot_plate
    from app.reporting.result_snapshot import ResultRow
    import matplotlib.pyplot as plt
    original = plt.subplots
    observed = []
    def subplots(**kwargs):
        figure, axes = original(**kwargs)
        observed.append(axes)
        return figure, axes
    monkeypatch.setattr(snapshot_plate.plt, "subplots", subplots)
    rows = [ResultRow("P24", "", "Unknown", None, None, None, 4, "missing", "missing")]
    assert snapshot_plate.render_snapshot_plate(rows).startswith(b"\x89PNG")
    assert observed[0].get_xlim()[1] > 24
    assert observed[0].get_ylim()[0] > 16


@pytest.mark.parametrize("extension", ["pdf", "xlsx"])
def test_multimarker_missing_rows_preserved(plate: SimpleNamespace, extension: str, monkeypatch: pytest.MonkeyPatch) -> None:
    from fixtures_ux_followup import make_ux_markers
    from app.reporting import snapshot_pdf, snapshot_xlsx
    unified = plate.upload.sessions["export"]
    unified.data = [reading for reading in unified.data if not (reading.well == "A1" and reading.cycle == 20)]
    for well, kind in (("A2", "Omit"), ("A3", "Empty"), ("A4", "Unknown")):
        assert plate.client.post("/api/data/export/welltypes", json={"wells": [well], "well_type": kind}).status_code == 200
    assert plate.client.post("/api/data/export/cluster", json={
        "cycle": 20, "use_rox": True, "regions": [marker.model_dump() for marker in make_ux_markers(unified)],
    }).status_code == 200
    module = snapshot_pdf if extension == "pdf" else snapshot_xlsx
    original = module.render_scatter_png
    ploidies = []
    def render(points, dye, **kwargs):
        ploidies.append(kwargs["ploidy"])
        assert kwargs["coordinate_basis"] == "reference-normalized / raw fallback"
        return original(points, dye, **kwargs)
    monkeypatch.setattr(module, "render_scatter_png", render)
    response = plate.client.get(f"/api/data/export/export/{extension}")
    assert response.status_code == 200
    assert set(ploidies) == {marker.ploidy for marker in make_ux_markers(unified)}
    if extension == "xlsx":
        workbook = load_workbook(io.BytesIO(response.content))
        headers = [cell.value for cell in workbook["Results"][1]]
        rows = {row[0]: dict(zip(headers, row)) for row in workbook["Results"].iter_rows(min_row=2, values_only=True)}
        assert len(rows) == 96
        assert rows["A1"]["Read Status"] == "missing" and rows["A1"]["FAM (norm)"] is None
        assert [rows[well]["Genotype"] for well in ("A2", "A3", "A4")] == ["Omit", "Empty", "Unknown"]
    else:
        text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
        for token in ("missing", "Omit", "Empty", "Unknown", "outside_marker", "H12", "Plate legend"):
            assert token in text


def _bind_report_asg() -> None:
    from app.asg_client import ASGLaunchContext, ASGLaunchSaveCredential
    from app.asg_session import bind_session_to_current_asg_launch, remember_asg_launch
    remember_asg_launch("u", ASGLaunchContext("design_run_item", "1", {"name": "accepted"}),
                        ASGLaunchSaveCredential("launch", "test-token"), ["snp:save_result"], None)
    bind_session_to_current_asg_launch("export", "u")


def test_xlsx_legacy_called_policy_preserves_empty_and_omit() -> None:
    from openpyxl import Workbook
    from app.reporting.snapshot_xlsx import _summary_qc
    sheet = Workbook().active
    _summary_qc(sheet, ["Empty", "Omit", "Allele 1 Homo", "Unknown", "NTC", "Undetermined", "Unassigned"])
    values = dict(sheet.values)
    assert values["Called"] == 3
    assert values["Call rate (%)"] == 42.9


def test_xlsx_rejects_oversized_derived_caption(plate: SimpleNamespace, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.reporting import snapshot_xlsx
    from app.reporting.snapshot_presentation import ReportFigure
    monkeypatch.setattr(snapshot_xlsx, "report_figures", lambda *_: [ReportFigure("x" * 33000, 2, [])])
    response = plate.client.get("/api/data/export/export/xlsx")
    assert response.status_code == 400


def test_xlsx_preserves_negative_numeric_signal(plate: SimpleNamespace) -> None:
    data = plate.upload.sessions["export"]
    for reading in data.data:
        if reading.well == "A1" and reading.cycle == 20:
            reading.fam = -12.5
    assert plate.client.post("/api/data/export/cluster", json={"cycle": 20, "use_rox": False}).status_code == 200
    response = plate.client.get("/api/data/export/export/xlsx")
    sheet = load_workbook(io.BytesIO(response.content))["Results"]
    headers = [cell.value for cell in sheet[1]]
    cell = sheet.cell(2, headers.index("FAM (raw)") + 1)
    assert cell.value == -12.5 and cell.data_type == "n"


@pytest.mark.parametrize("extension", ["pdf", "xlsx"])
def test_report_plot_failure_does_not_return_partial_success(plate: SimpleNamespace, monkeypatch: pytest.MonkeyPatch, extension: str) -> None:
    from app.reporting import snapshot_pdf, snapshot_xlsx
    module = snapshot_pdf if extension == "pdf" else snapshot_xlsx
    def fail(*args, **kwargs):
        raise RuntimeError("plot failed")
    monkeypatch.setattr(module, "render_scatter_png", fail)
    with pytest.raises(RuntimeError, match="plot failed"):
        plate.client.get(f"/api/data/export/export/{extension}")


def test_pdf_paired_call_confidence_and_background_numeric(plate: SimpleNamespace, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.reporting import snapshot_pdf
    from app.routers.sample import sample_name_store
    sample_name_store["export"] = {"A1": "PAIR_A1", "A2": "PAIR_A2"}
    assert plate.client.post("/api/data/export/cluster", json={"cycle": 20, "use_rox": True, "background": "pre_read"}).status_code == 200
    result = plate.clustering.cluster_store["export"]
    captured = []
    # The threshold fixture has no confidence output; seed a saved adapter value.
    result.confidences = {"A1": 0.873}
    original = snapshot_pdf.render_scatter_png
    def render(points, dye, **kwargs):
        captured.extend(points)
        return original(points, dye, **kwargs)
    monkeypatch.setattr(snapshot_pdf, "render_scatter_png", render)
    response = plate.client.get("/api/data/export/export/pdf")
    assert captured[0]["norm_fam"] == 8.0
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    # Each identity row puts the sample, saved label and confidence together.
    start = text.index("PAIR_A1")
    end = text.index("PAIR_A2", start)
    row_text = text[start:end]
    assert result.assignments["A1"] in row_text
    assert str(round(result.confidences["A1"] * 100, 1)) in row_text


@pytest.mark.parametrize("state,code", [
    ("missing", "NO_COMPLETED_RESULT"), ("legacy", "LEGACY_CONTEXT_UNKNOWN"),
    ("stale", "INPUT_REVISION_CONFLICT"), ("mismatch", "EXPORT_CONDITION_MISMATCH"),
    ("pending", "ANALYSIS_IN_PROGRESS"), ("replaced", "RESULT_REVISION_CONFLICT"),
])
def test_asg_rejection_never_posts(plate: SimpleNamespace, monkeypatch: pytest.MonkeyPatch, state: str, code: str) -> None:
    from unittest.mock import Mock
    from app.routers import asg
    from app.models import ClusteringRequest
    from app.processing.analysis_state import fail_analysis
    _bind_report_asg()
    monkeypatch.setattr(asg, "is_asg_launch_mode", lambda: True)
    post = Mock()
    monkeypatch.setattr(asg, "post_analysis_result", post)
    body = {"session_id": "export"}
    ticket = None
    if state == "missing":
        plate.clustering.cluster_store.pop("export")
    elif state == "legacy":
        plate.clustering.cluster_store["export"].analysis_context = None
    elif state == "stale":
        plate.upload.sessions["export"].input_revision += 1
    elif state == "pending":
        ticket, _ = plate.clustering._capture_analysis("export", ClusteringRequest(cycle=20))
    elif state == "replaced":
        body["result_revision"] = "00000000-0000-0000-0000-000000000000"
    else:
        body["selected_cycle"] = 40
    try:
        response = plate.client.post("/api/asg/save-result", json=body)
        assert response.status_code == 409
        assert response.json()["detail"]["code"] == code
        post.assert_not_called()
    finally:
        if ticket:
            fail_analysis(ticket)


def test_asg_schema2_freezes_launch_and_manual_unknown(plate: SimpleNamespace, monkeypatch: pytest.MonkeyPatch) -> None:
    from app import asg_result
    from app.asg_session import get_session_asg_launch
    from app.auth import TokenData
    _bind_report_asg()
    assert plate.client.post("/api/data/export/welltypes", json={"wells": ["A1"], "well_type": "Unknown"}).status_code == 200
    assert plate.client.post("/api/data/export/cluster", json={"cycle": 20, "ploidy": 4}).status_code == 200
    original = asg_result._render_asg
    def render(snapshot, launch):
        get_session_asg_launch("export").context["name"] = "later"
        plate.clustering.welltype_store["export"]["A1"] = "NTC"
        return original(snapshot, launch)
    monkeypatch.setattr(asg_result, "_render_asg", render)
    payload = asg_result.build_result_snapshot("export", user=TokenData(user_id="u", username="u", role="user"))
    assert payload["schema_version"] == 2
    assert payload["summary"]["allele_frequency"] is None and payload["summary"]["hwe"] is None
    assert payload["result"]["wells"][0]["effective_type"] == "Unknown"
    assert payload["result"]["asg_target"]["context"]["name"] == "accepted"
