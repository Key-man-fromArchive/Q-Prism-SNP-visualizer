"""Marker figure titles come from one function; single-marker downloads are named by marker."""
import io
import re
import zipfile

from pptx import Presentation

from test_export_pptx import run as _run, data_client as _data_client  # noqa: F401

run = _run
data_client = _data_client
BASE = "/api/data/{sid}/export/"


def _disposition(response) -> str:
    return response.headers["content-disposition"]


def _plain_name(response) -> str:
    return _disposition(response).split("filename=")[1].split(";")[0].strip('"')


def _first_marker(run):
    return run.markers[0]


def test_zip_chart_titles_name_marker_and_cycle_without_internal_ids(run, monkeypatch) -> None:
    from app.reporting import snapshot_images
    titles: list[str] = []
    real = snapshot_images.render_scatter_png

    def spy(*args, **kwargs):
        titles.append(kwargs["title"])
        return real(*args, **kwargs)

    monkeypatch.setattr(snapshot_images, "render_scatter_png", spy)
    assert run.client.get(BASE.format(sid=run.sid) + "scatter-png.zip").status_code == 200
    assert len(titles) == 6
    assert titles[0].startswith("QPrism1 · ")
    for title, marker in zip(titles, run.markers, strict=True):
        assert "[" not in title and "ploidy" not in title and marker["id"] not in title


def test_pptx_slide_title_matches_shared_title_and_hides_id(run) -> None:
    response = run.client.get(BASE.format(sid=run.sid) + "pptx")
    deck = Presentation(io.BytesIO(response.content))
    texts = [shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
    assert any(t.startswith("QPrism1 · ") for t in texts)
    assert not any(re.search(r"\[[^\]]+\] / ploidy", t) for t in texts)


def test_single_marker_pdf_and_zip_filenames_follow_pptx_rule(run) -> None:
    marker = _first_marker(run)
    pptx = _plain_name(run.client.get(BASE.format(sid=run.sid) + "pptx", params={"marker_ids": marker["id"]}))
    cycle = re.search(r"_cycle(\d+)\.pptx$", pptx).group(1)
    pdf = run.client.get(BASE.format(sid=run.sid) + "pdf", params={"marker_ids": marker["id"]})
    zipped = run.client.get(BASE.format(sid=run.sid) + "scatter-png.zip", params={"marker_ids": marker["id"]})
    assert _plain_name(pdf) == f"snp_report_{marker['name']}_cycle{cycle}.pdf"
    assert _plain_name(zipped) == f"snp_scatter_png_{marker['name']}_cycle{cycle}.zip"
    assert "filename*=UTF-8''" in _disposition(pdf)
    assert "filename*=UTF-8''" in _disposition(zipped)


def test_multi_marker_and_whole_run_filenames_unchanged(run) -> None:
    two = ",".join(m["id"] for m in run.markers[:2])
    assert "whole-run" in _plain_name(run.client.get(BASE.format(sid=run.sid) + "pdf"))
    assert re.fullmatch(r"snp_scatter_png_cycle\d+\.zip",
                        _plain_name(run.client.get(BASE.format(sid=run.sid) + "scatter-png.zip")))
    names = _plain_name(run.client.get(BASE.format(sid=run.sid) + "pdf", params={"marker_ids": two}))
    assert names.startswith("snp_report_") and names.endswith(".pdf")
    assert zipfile.is_zipfile(io.BytesIO(run.client.get(
        BASE.format(sid=run.sid) + "scatter-png.zip", params={"marker_ids": two}).content))


def test_hostile_marker_name_stays_safe_in_filename(run) -> None:
    marker = _first_marker(run)
    run.client.put(f"/api/data/{run.sid}/markers/{marker['id']}", json={"name": "../$마커\x01:x"})
    response = run.client.get(BASE.format(sid=run.sid) + "pdf", params={"marker_ids": marker["id"]})
    assert response.status_code == 200
    plain = _plain_name(response)
    assert "/" not in plain and ".." not in plain and "\x01" not in _disposition(response)
