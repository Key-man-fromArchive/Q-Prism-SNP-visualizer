"""Large tables are read row by row and stop at the row limit."""
from __future__ import annotations

import re
import time
import tracemalloc
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.import_errors import ImportErrorCode, ImportValidationError
from app.parsers import generic_table
from app.parsers.generic_table import (
    MAX_IMPORT_COLUMNS,
    MAX_IMPORT_ROWS,
    _minimal_config,
    _read_table,
)
from app.parsers.registry import build_default_parser_registry

FIXTURES = Path(__file__).parent / "fixtures" / "import"


def _limit_codes(exc_info) -> set[str]:
    return _limit_codes_of(exc_info.value)


def _limit_codes_of(error: ImportValidationError) -> set[str]:
    return {issue.code for issue in error.issues}


def _measure(func):
    tracemalloc.start()
    started = time.perf_counter()
    try:
        outcome = None
        error = None
        try:
            outcome = func()
        except ImportValidationError as exc:
            error = exc
        elapsed = time.perf_counter() - started
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    return outcome, error, elapsed, peak


def _declared_wide_workbook(path: Path, *, data_row: int = 20000) -> Path:
    base = path.with_name("base.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["well", "cycle", "rfu"])
    wb.save(base)
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r'<dimension ref="[^"]*"\s*/>', '<dimension ref="A1:XFD1048576"/>', xml)
                row = f'<row r="{data_row}"><c r="A{data_row}" t="inlineStr"><is><t>A1</t></is></c></row>'
                xml = xml.replace("</sheetData>", row + "</sheetData>")
                data = xml.encode("utf-8")
            dst.writestr(item.filename, data)
    return path


def test_declared_dimension_does_not_size_the_read(tmp_path):
    path = _declared_wide_workbook(tmp_path / "declared.xlsx")
    _, error, elapsed, peak = _measure(lambda: _read_table(path, _minimal_config()))
    # Only 20000 rows are actually present, so the table is read and bounded by content.
    assert error is None
    assert elapsed < 2.0
    assert peak < 200 * 1024 * 1024


def test_declared_dimension_with_rows_past_the_limit_is_refused_quickly(tmp_path):
    path = _declared_wide_workbook(tmp_path / "declared.xlsx", data_row=MAX_IMPORT_ROWS + 500)
    _, error, elapsed, peak = _measure(lambda: _read_table(path, _minimal_config()))
    assert error is not None
    assert ImportErrorCode.FILE_LIMIT_EXCEEDED in _limit_codes_of(error)
    assert elapsed < 2.0
    assert peak < 200 * 1024 * 1024


def test_registry_sniff_of_declared_dimension_file_is_bounded(tmp_path):
    path = _declared_wide_workbook(tmp_path / "declared.xlsx", data_row=MAX_IMPORT_ROWS + 500)
    registry = build_default_parser_registry()
    matched, _, elapsed, peak = _measure(lambda: registry.match(path, "declared.xlsx"))
    assert matched is None
    assert elapsed < 2.0
    assert peak < 200 * 1024 * 1024


def test_blank_line_csv_is_refused_within_budget(tmp_path):
    path = tmp_path / "blank.csv"
    path.write_bytes(b"\n" * (20 * 1024 * 1024))
    _, error, elapsed, peak = _measure(lambda: _read_table(path, _minimal_config()))
    assert error is not None
    assert ImportErrorCode.FILE_LIMIT_EXCEEDED in _limit_codes_of(error)
    assert elapsed < 2.0
    assert peak < 300 * 1024 * 1024


def test_blank_line_tsv_registry_pass_is_bounded(tmp_path):
    path = tmp_path / "blank.tsv"
    path.write_bytes(b"\n" * (20 * 1024 * 1024))
    registry = build_default_parser_registry()
    matched, _, elapsed, peak = _measure(lambda: registry.match(path, "blank.tsv"))
    assert matched is None
    assert elapsed < 2.0
    assert peak < 300 * 1024 * 1024


def test_single_overlong_line_is_refused(tmp_path):
    path = tmp_path / "long.csv"
    path.write_text("well,cycle,rfu\n" + "x" * (generic_table.MAX_IMPORT_LINE_CHARS + 10) + "\n")
    with pytest.raises(ImportValidationError) as exc_info:
        _read_table(path, _minimal_config())
    assert ImportErrorCode.FILE_LIMIT_EXCEEDED in _limit_codes(exc_info)


def test_too_many_columns_is_refused(tmp_path):
    path = tmp_path / "wide.csv"
    path.write_text(",".join(f"c{i}" for i in range(MAX_IMPORT_COLUMNS + 1)) + "\n1\n")
    with pytest.raises(ImportValidationError) as exc_info:
        _read_table(path, _minimal_config())
    assert ImportErrorCode.FILE_LIMIT_EXCEEDED in _limit_codes(exc_info)


@pytest.mark.parametrize("suffix", [".csv", ".tsv"])
def test_rows_beyond_the_limit_give_limit_error(tmp_path, suffix):
    delimiter = "," if suffix == ".csv" else "\t"
    path = tmp_path / f"many{suffix}"
    lines = [delimiter.join(["well", "cycle", "rfu"])]
    lines += [delimiter.join(["A1", str(i), "1.0"]) for i in range(MAX_IMPORT_ROWS + 5)]
    path.write_text("\n".join(lines) + "\n")
    with pytest.raises(ImportValidationError) as exc_info:
        _read_table(path, _minimal_config())
    assert ImportErrorCode.FILE_LIMIT_EXCEEDED in _limit_codes(exc_info)


def test_xlsx_rows_beyond_the_limit_give_limit_error(tmp_path):
    path = tmp_path / "many.xlsx"
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["well", "cycle", "rfu"])
    for i in range(MAX_IMPORT_ROWS + 5):
        ws.append(["A1", i, 1.0])
    wb.save(path)
    with pytest.raises(ImportValidationError) as exc_info:
        _read_table(path, _minimal_config())
    assert ImportErrorCode.FILE_LIMIT_EXCEEDED in _limit_codes(exc_info)


def test_rows_at_the_limit_are_accepted(tmp_path):
    path = tmp_path / "edge.csv"
    lines = ["well,cycle,rfu"] + [f"A1,{i},1.0" for i in range(MAX_IMPORT_ROWS)]
    path.write_text("\n".join(lines) + "\n")
    table = _read_table(path, _minimal_config())
    assert len(table.rows) == MAX_IMPORT_ROWS


def test_xlsx_trailing_empty_cells_are_trimmed(tmp_path):
    path = tmp_path / "small.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["well", "cycle", "rfu"])
    ws.append(["A1", 1, 100.5])
    wb.save(path)
    table = _read_table(path, _minimal_config())
    assert table.headers == ["well", "cycle", "rfu"]
    assert table.rows == [{"well": "A1", "cycle": "1", "rfu": "100.5"}]


def _fixture_files() -> list[Path]:
    return sorted(
        p for p in FIXTURES.rglob("*")
        if p.is_file() and p.suffix.lower() in {".csv", ".tsv", ".txt", ".xlsx"}
    )


def test_fixtures_read_the_same_with_and_without_a_shared_scope():
    files = _fixture_files()
    assert files
    for path in files:
        def read():
            try:
                table = _read_table(path, _minimal_config())
            except ImportValidationError as exc:
                return ("error", tuple(issue.code for issue in exc.issues))
            return (table.headers, table.rows, table.delimiter)

        plain = read()
        with generic_table.table_read_scope():
            first = read()
            second = read()
        assert plain == first == second


@pytest.mark.parametrize("kind", ["tsv", "csv", "xlsx"])
def test_sniff_pass_reads_each_file_once(tmp_path, monkeypatch, kind):
    if kind == "xlsx":
        path = tmp_path / "plain.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["a", "b"])
        wb.active.append(["1", "2"])
        wb.save(path)
    else:
        delimiter = "\t" if kind == "tsv" else ","
        path = tmp_path / f"plain.{kind}"
        path.write_text("a{d}b\n1{d}2\n".format(d=delimiter))

    calls = {"delimited": 0, "xlsx": 0}
    real_delimited = generic_table._load_delimited_matrix
    real_xlsx = generic_table._load_xlsx_matrix

    def spy_delimited(*args, **kwargs):
        calls["delimited"] += 1
        return real_delimited(*args, **kwargs)

    def spy_xlsx(*args, **kwargs):
        calls["xlsx"] += 1
        return real_xlsx(*args, **kwargs)

    monkeypatch.setattr(generic_table, "_load_delimited_matrix", spy_delimited)
    monkeypatch.setattr(generic_table, "_load_xlsx_matrix", spy_xlsx)

    build_default_parser_registry().match(path, path.name)

    assert calls == {"delimited": 0 if kind == "xlsx" else 1, "xlsx": 1 if kind == "xlsx" else 0}


def test_xlsx_sparse_rows_and_columns_keep_their_positions(tmp_path):
    path = tmp_path / "sparse.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "well"
    ws["C1"] = "rfu"
    ws["A4"] = "B2"
    ws["C4"] = 12.5
    wb.save(path)
    table = _read_table(path, _minimal_config())
    assert table.headers == ["well", "", "rfu"]
    assert [row for _, row in table.iter_data_rows()][-1] == {"well": "B2", "rfu": "12.5"}
    assert len(table.rows) == 3
