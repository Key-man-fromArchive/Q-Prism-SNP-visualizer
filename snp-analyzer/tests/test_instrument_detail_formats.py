"""P7-B2: every supported format reports the instrument identity its file states."""
import io
import os
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.import_models import AssayModeId, ImportRole, MappingConfig, NormalizationMode
from app.models import InstrumentDetail
from app.parsers import cfx_opus, cfx_xml_parser, pcrd_raw, quantstudio
from app.parsers.detector import detect_and_parse
from app.parsers.instrument_detail import (
    eds_instrument_detail,
    messages_log_model,
    normalize_cfx_model,
    pcrd_instrument_detail,
    rdml_instrument_detail,
    xls_instrument_detail,
)
from app.parsers.rdml import RDMLParser
from tests.quantstudio_fixtures import build_quantstudio_eds

BIORAD = "Bio-Rad"
AB = "Applied Biosystems"


# --- CFX .pcrd ---------------------------------------------------------

@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("CFX Duet", "CFX Duet"),
        ("CFX Opus 96", "CFX Opus 96"),
        ("Opus 96", "CFX Opus 96"),
        ("  Opus 96 ", "CFX Opus 96"),
        (None, None),
        ("", None),
    ],
)
def test_normalize_cfx_model(raw, expected):
    assert normalize_cfx_model(raw) == expected


def _pcrd_root(entry=None, version=None):
    root = ET.Element("experimentalData2")
    header = ET.SubElement(root, "header")
    if version is not None:
        header.set("createdByClientAppVersion", version)
    if entry is not None:
        ET.SubElement(root, "machineProperties", machinePropertiesEntry=entry)
    return root


def test_pcrd_detail_from_machine_properties_and_header():
    root = _pcrd_root(
        "Serial Number : 123|Block Description : Opus 96|Lid : 105",
        "5.3.022.1030 (Windows NT 10.0)",
    )
    assert pcrd_instrument_detail(root) == InstrumentDetail(
        vendor=BIORAD, model="CFX Opus 96", software="CFX Maestro 5.3.022.1030"
    )


def test_pcrd_duet_and_missing_pieces():
    duet = pcrd_instrument_detail(_pcrd_root("Block Description : CFX Duet"))
    assert (duet.model, duet.software) == ("CFX Duet", None)
    only_version = pcrd_instrument_detail(_pcrd_root(None, "5.3.022.1030"))
    assert (only_version.model, only_version.software) == (None, "CFX Maestro 5.3.022.1030")
    assert pcrd_instrument_detail(_pcrd_root()) is None


def _minimal_pcrd_root(entry):
    root = _pcrd_root(entry, "5.3.022.1030 (Windows)")
    setup = ET.SubElement(root, "plateSetup2", rows="8", columns="12")
    layers = ET.SubElement(setup, "dyeLayersList")
    for name, channel in (("FAM", "0"), ("VIC", "1")):
        layer = ET.SubElement(layers, "dyeLayer", plateName=name)
        ET.SubElement(layer, "fluor", channelPosition=channel)
        wells = ET.SubElement(layer, "wellSamples")
        ET.SubElement(wells, "wellSample", plateIndex="0", wellSampleType="wcSample", sampleId="S1")
    read = ET.SubElement(ET.SubElement(root, "runData"), "plateReadDataVector")
    inner = ET.SubElement(ET.SubElement(read, "plateRead"), "PlateRead")
    hdr = ET.SubElement(ET.SubElement(inner, "Hdr"), "PlateReadDataHeader")
    for tag, value in (("Step", "1"), ("Cycle", "1"), ("ChCount", "2"), ("NumRows", "9"), ("NumCols", "12")):
        ET.SubElement(hdr, tag).text = value
    ET.SubElement(ET.SubElement(inner, "Data"), "PAr").text = ";".join(["100"] * (2 * 108 * 4))
    return root


def test_parse_pcrd_uses_model_for_instrument_string(monkeypatch):
    monkeypatch.setattr(pcrd_raw, "_extract_xml", lambda _p: _minimal_pcrd_root("Block Description : CFX Duet"))
    unified = pcrd_raw.parse_pcrd("ignored.pcrd")
    assert unified.instrument == "CFX Duet (raw)"
    assert unified.instrument_detail.software == "CFX Maestro 5.3.022.1030"
    monkeypatch.setattr(pcrd_raw, "_extract_xml", lambda _p: _minimal_pcrd_root(None))
    assert pcrd_raw.parse_pcrd("ignored.pcrd").instrument == "Bio-Rad CFX (raw)"


# --- QuantStudio .eds --------------------------------------------------

LOG = (
    "noise line\n"
    "Instrument Properties: -serial=1 -product=QuantStudio3_5 -x=2\n"
    "Instrument Properties: -product=Other\n"
)


@pytest.mark.parametrize(
    ("product", "model"),
    [
        ("QuantStudio3_5", "QuantStudio 3/5"),
        ("QuantStudio5", "QuantStudio 5"),
        ("QuantStudio12K", "QuantStudio 12K"),
    ],
)
def test_messages_log_model(product, model):
    line = f"Instrument Properties: -a=1 -product={product} -b=2\n"
    assert messages_log_model(io.BytesIO(line.encode())) == model


def test_messages_log_model_stops_at_first_match_and_limit():
    assert messages_log_model(io.BytesIO(LOG.encode())) == "QuantStudio 3/5"
    assert messages_log_model(io.BytesIO(b"nothing here\n")) is None
    padded = b"x" * 100 + b"\n"
    late = padded * 50 + LOG.encode()
    assert messages_log_model(io.BytesIO(late), limit=1000) is None


def test_eds_software_joins_title_and_version_for_quantstudio():
    manifest = (
        b"Implementation-Title: QuantStudio 3 and 5 Software\r\n"
        b"Implementation-Version: 1.5.3\r\n"
    )
    exp = b"<Experiment><InstrumentTypeId>quantstudio3</InstrumentTypeId></Experiment>"
    detail = eds_instrument_detail(exp, manifest, log_model="QuantStudio 3/5")
    assert detail == InstrumentDetail(
        vendor=AB, model="QuantStudio 3/5", software="QuantStudio 3 and 5 Software 1.5.3"
    )


def test_eds_log_model_alone_names_the_vendor():
    detail = eds_instrument_detail(None, None, log_model="QuantStudio 3/5")
    assert (detail.vendor, detail.model, detail.software) == (AB, "QuantStudio 3/5", None)


def _eds_with(members):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(build_quantstudio_eds())) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            if item.filename not in members:
                dst.writestr(item.filename, src.read(item.filename))
        for name, text in members.items():
            dst.writestr(name, text)
    return out.getvalue()


def test_quantstudio_eds_instrument_string_follows_model(tmp_path):
    raw = _eds_with({"apldbio/sds/messages.log": LOG})
    path = tmp_path / "run.eds"
    path.write_bytes(raw)
    unified = detect_and_parse(str(path), path.name)
    assert unified.instrument_detail.model == "QuantStudio 3/5"
    assert unified.instrument == "QuantStudio 3/5 (raw)"


# --- QuantStudio / StepOne .xls export --------------------------------

@pytest.mark.parametrize(
    ("value", "model"),
    [
        ("steponeplus", "StepOnePlus"),
        ("stepone", "StepOne"),
        ("quantstudio3", "QuantStudio 3"),
        ("QuantStudio 5", "QuantStudio 5"),
        ("viia7", "viia7"),
    ],
)
def test_xls_detail_models(value, model):
    rows = [["File Name", "a.eds"], ["Instrument Type", value], ["Well", "Cycle"]]
    assert xls_instrument_detail(rows) == InstrumentDetail(vendor=AB, model=model, software=None)


def test_xls_without_instrument_type_is_none():
    assert xls_instrument_detail([["File Name", "a.eds"]]) is None
    assert xls_instrument_detail([["Instrument Type", ""]]) is None


def _xls_rows_sheet(instrument_type):
    """Duck-typed stand-in for an xlrd sheet."""
    class Sheet:
        def __init__(self, rows):
            self.rows = rows
            self.nrows = len(rows)
            self.ncols = max(len(r) for r in rows)

        def cell_value(self, r, c):
            row = self.rows[r]
            return row[c] if c < len(row) else ""

    rows = [["Well", "Cycle", "VIC", "ROX", "FAM"], [1, 1, 10.0, 5.0, 20.0]]
    if instrument_type is not None:
        rows.insert(0, ["Instrument Type", instrument_type])
    return Sheet(rows)


class _Book:
    def __init__(self, sheet):
        self.sheet = sheet

    def sheet_by_index(self, _i):
        return self.sheet


@pytest.mark.parametrize(
    ("itype", "instrument", "detail_model"),
    [
        ("steponeplus", "StepOnePlus", "StepOnePlus"),
        (None, "QuantStudio 3", None),
    ],
)
def test_parse_quantstudio_uses_sheet_instrument_type(monkeypatch, itype, instrument, detail_model):
    monkeypatch.setattr(
        quantstudio.xlrd, "open_workbook", lambda _p: _Book(_xls_rows_sheet(itype))
    )
    unified = quantstudio.parse_quantstudio("x.xls")
    assert unified.instrument == instrument
    assert (unified.instrument_detail.model if unified.instrument_detail else None) == detail_model


# --- CFX xlsx / xml / zip exports --------------------------------------

def _cfx_workbook(version):
    wb = openpyxl.Workbook()
    fam = wb.active
    fam.title = "FAM"
    hex_ = wb.create_sheet("HEX")
    for ws in (fam, hex_):
        ws.append([None, "Cycle", "A1"])
        ws.append([None, 1, 1000.0])
    info = wb.create_sheet("Run Information")
    info.append(["File Name", "run.pcrd"])
    if version is not None:
        info.append(["CFX Maestro Version", version])
    return wb


@pytest.mark.parametrize("version", ["5.3.022.1030", "CFX Maestro 5.3.022.1030"])
def test_cfx_xlsx_detail_from_run_information(version):
    unified = cfx_opus._parse_workbook(_cfx_workbook(version))
    assert unified.instrument == "Bio-Rad CFX"
    assert unified.instrument_detail == InstrumentDetail(
        vendor=BIORAD, model=None, software="CFX Maestro 5.3.022.1030"
    )


def test_cfx_xlsx_without_version_has_no_detail():
    unified = cfx_opus._parse_workbook(_cfx_workbook(None))
    assert unified.instrument_detail is None
    assert unified.instrument == "Bio-Rad CFX"


ADSHEET = (
    "<ADSheet><Row><Well>A01</Well><Sample>S1</Sample><Call>Allele 1</Call>"
    "<Type>Unknown</Type><RFU1>100</RFU1><RFU2>50</RFU2></Row></ADSheet>"
)


def _zip_with(tmp_path, run_info_name, run_info_xml):
    path = tmp_path / "export.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("run - Allelic Discrimination Results_ADSheet.xml", ADSHEET)
        if run_info_xml is not None:
            zf.writestr(run_info_name, run_info_xml)
    return str(path)


@pytest.mark.parametrize(
    "run_info",
    [
        "<Run_x0020_Information><Row><CFX_x0020_Maestro_x0020_Version>5.3.022.1030"
        "</CFX_x0020_Maestro_x0020_Version></Row></Run_x0020_Information>",
        "<Run_x0020_Information><Row><Key>CFX Maestro Version</Key>"
        "<Value>5.3.022.1030</Value></Row></Run_x0020_Information>",
    ],
)
def test_cfx_zip_detail_from_run_information(tmp_path, run_info):
    path = _zip_with(tmp_path, "run - Run Information.xml", run_info)
    unified = cfx_xml_parser.parse_cfx_xml_zip(path)
    assert unified.instrument_detail == InstrumentDetail(
        vendor=BIORAD, model=None, software="CFX Maestro 5.3.022.1030"
    )
    assert unified.instrument == "Bio-Rad CFX"


def test_cfx_zip_without_run_information_has_no_detail(tmp_path):
    unified = cfx_xml_parser.parse_cfx_xml_zip(_zip_with(tmp_path, "", None))
    assert unified.instrument_detail is None
    assert unified.instrument == "Bio-Rad CFX"


# --- RDML ---------------------------------------------------------------

RDML_FIXTURE = Path(__file__).parent / "fixtures" / "import" / "rdml" / "wt_mt.rdml"


def _rdml_unified(tmp_path, extra):
    text = RDML_FIXTURE.read_text(encoding="utf-8").replace(
        '<rdml version="1.3">', '<rdml version="1.3">' + extra, 1
    )
    path = tmp_path / "a.rdml"
    path.write_text(text, encoding="utf-8")
    mapping = MappingConfig(
        assay_mode=AssayModeId.WT_MT,
        normalization_mode=NormalizationMode.NONE,
        channel_roles={"FAM": ImportRole.WT, "VIC": ImportRole.MT1},
        rdml_run_id="run-1",
        rdml_target_ids=["WT", "MT"],
    )
    parser = RDMLParser()
    return parser.to_unified(parser.parse(path, path.name, mapping))


def test_rdml_detail_from_instrument_vendor_and_software(tmp_path):
    unified = _rdml_unified(
        tmp_path,
        "<instrument>CFX96</instrument><manufacturer>Bio-Rad</manufacturer>"
        "<dataCollectionSoftware><name>CFX Maestro</name><version>5.3</version>"
        "</dataCollectionSoftware>",
    )
    assert unified.instrument_detail == InstrumentDetail(
        vendor="Bio-Rad", model="CFX96", software="CFX Maestro 5.3"
    )


def test_rdml_detail_falls_back_to_documentation_software(tmp_path):
    unified = _rdml_unified(tmp_path, "")
    assert unified.instrument_detail == InstrumentDetail(
        vendor=None, model=None, software="Q-Prism Synthetic RDML"
    )
    assert rdml_instrument_detail(None, None, None) is None


# --- real files (optional) ---------------------------------------------

def _real(name):
    value = os.environ.get(name)
    return value if value and Path(value).is_file() else None


@pytest.mark.skipif(not _real("QPRISM_SAMPLE_PCRD_DUET"), reason="QPRISM_SAMPLE_PCRD_DUET not set")
def test_real_duet_pcrd():
    path = _real("QPRISM_SAMPLE_PCRD_DUET")
    detail = detect_and_parse(path, Path(path).name).instrument_detail
    assert (detail.vendor, detail.model) == (BIORAD, "CFX Duet")
    assert detail.software.startswith("CFX Maestro ")


@pytest.mark.skipif(not _real("QPRISM_SAMPLE_PCRD_OPUS"), reason="QPRISM_SAMPLE_PCRD_OPUS not set")
def test_real_opus_pcrd():
    path = _real("QPRISM_SAMPLE_PCRD_OPUS")
    detail = detect_and_parse(path, Path(path).name).instrument_detail
    assert detail.vendor == BIORAD
    assert detail.model in {"CFX Opus 96", "CFX Opus"} or detail.model.startswith("CFX Opus")
    assert detail.software.startswith("CFX Maestro ")


@pytest.mark.skipif(not _real("QPRISM_SAMPLE_QS_EDS"), reason="QPRISM_SAMPLE_QS_EDS not set")
def test_real_quantstudio_eds():
    path = _real("QPRISM_SAMPLE_QS_EDS")
    unified = detect_and_parse(path, Path(path).name)
    detail = unified.instrument_detail
    assert (detail.vendor, detail.model) == (AB, "QuantStudio 3/5")
    assert detail.software == "QuantStudio 3 and 5 Software 1.5.3"
    assert unified.instrument.startswith("QuantStudio 3/5")
