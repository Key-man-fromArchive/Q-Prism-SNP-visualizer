"""Instrument files are read within the sizes the instruments write.

Covers per-member size limits, refusal of documents with a DTD (in any text
encoding), plate and read count limits, the archive directory check that runs
before an archive is opened, and the single streaming read of CFX workbooks.
"""

from __future__ import annotations

import io
import os
import struct
import time
import zipfile

import openpyxl
import pytest

from app.import_errors import ImportValidationError
from app.parsers import cfx_opus, cfx_xml_parser, detector, pcrd_raw, safe_xml
from app.parsers.detector import detect_and_parse
from app.parsers.eds_common import _parse_plate_dims
from app.parsers.eds_raw import _parse_multicomponent, parse_eds
from app.parsers.rdml import RDMLParser
from app.parsers.stepone_eds import plan_reads

from tests.quantstudio_fixtures import QuantStudioOptions, build_quantstudio_eds

DOCTYPE_ENTITY = (
    '<!DOCTYPE doc [<!ENTITY x "expanded">]><{root}><v>&x;</v></{root}>'
)
DOCTYPE_ONLY = "<!DOCTYPE {root} SYSTEM \"file.dtd\"><{root}/>"


def _utf8(text: str) -> bytes:
    return ('<?xml version="1.0" encoding="UTF-8"?>' + text).encode("utf-8")


def _utf16(text: str) -> bytes:
    return ('<?xml version="1.0" encoding="UTF-16"?>' + text).encode("utf-16")


ENCODERS = pytest.mark.parametrize("encode", [_utf8, _utf16], ids=["utf8", "utf16"])
DECLARATIONS = pytest.mark.parametrize(
    "template", [DOCTYPE_ENTITY, DOCTYPE_ONLY], ids=["entity", "doctype"]
)


def _replace_member(eds: bytes, name: str, payload: bytes, stored: bool = False) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(eds)) as src, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for info in src.infolist():
            if info.filename.endswith(name):
                method = zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED
                dst.writestr(info.filename, payload, compress_type=method)
            else:
                dst.writestr(info.filename, src.read(info))
    return out.getvalue()


def _write(tmp_path, name: str, data: bytes) -> str:
    path = tmp_path / name
    path.write_bytes(data)
    return str(path)


# --- size limits ---------------------------------------------------------


def test_oversize_xml_member_is_refused_quickly(tmp_path):
    big = _replace_member(
        build_quantstudio_eds(),
        "multicomponentdata.xml",
        os.urandom(26 * 1024 * 1024),
        stored=True,
    )
    path = _write(tmp_path, "big.eds", big)

    started = time.perf_counter()
    with pytest.raises(ValueError, match="larger than an instrument file can be"):
        detect_and_parse(path, "big.eds")
    assert time.perf_counter() - started < 5


def test_member_cap_applies_to_each_xml_document(tmp_path, monkeypatch):
    monkeypatch.setattr(safe_xml, "MAX_XML_BYTES", 1000)
    path = _write(tmp_path, "ok.eds", build_quantstudio_eds())
    with pytest.raises(ValueError, match="larger than an instrument file can be"):
        parse_eds(path)


def test_member_within_the_cap_still_parses(tmp_path):
    path = _write(tmp_path, "ok.eds", build_quantstudio_eds())
    assert detect_and_parse(path, "ok.eds").wells


# --- documents with a DTD ------------------------------------------------


@ENCODERS
@DECLARATIONS
@pytest.mark.parametrize(
    "member",
    ["multicomponentdata.xml", "plate_setup.xml", "tcprotocol.xml", "experiment.xml"],
)
def test_eds_documents_with_a_dtd_are_refused(tmp_path, encode, template, member):
    payload = encode(template.format(root="Doc"))
    path = _write(
        tmp_path, "dtd.eds", _replace_member(build_quantstudio_eds(), member, payload)
    )
    with pytest.raises(ValueError, match="document type declaration"):
        detect_and_parse(path, "dtd.eds")


@ENCODERS
@DECLARATIONS
def test_rdml_documents_with_a_dtd_are_refused(tmp_path, encode, template):
    path = tmp_path / "dtd.rdml"
    path.write_bytes(encode(template.format(root="rdml")))
    with pytest.raises(ImportValidationError) as exc_info:
        RDMLParser().preview(path, path.name)
    issue = exc_info.value.issues[0]
    assert issue.code == "unsupported_content"
    assert "DTD or entity declarations" in issue.message


def test_rdml_malformed_xml_gets_a_fixed_message(tmp_path):
    path = tmp_path / "bad.rdml"
    path.write_bytes(b"<rdml><unclosed></rdml>")
    with pytest.raises(ImportValidationError) as exc_info:
        RDMLParser().preview(path, path.name)
    assert exc_info.value.issues[0].message == "Malformed RDML XML."


@ENCODERS
def test_cfx_zip_documents_with_a_dtd_are_refused(tmp_path, encode):
    path = tmp_path / "export.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(
            "run - Allelic Discrimination Results_ADSheet.xml",
            encode(DOCTYPE_ENTITY.format(root="ADSheet")),
        )
    with pytest.raises(ValueError, match="document type declaration"):
        detect_and_parse(str(path), "export.zip")


def test_cfx_zip_xml_member_over_the_cap_is_refused(tmp_path, monkeypatch):
    monkeypatch.setattr(safe_xml, "MAX_XML_BYTES", 100)
    path = tmp_path / "export.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(
            "run - Allelic Discrimination Results_ADSheet.xml",
            "<ADSheet>" + "<Row/>" * 100 + "</ADSheet>",
        )
    with pytest.raises(ValueError, match="larger than an instrument file can be"):
        detect_and_parse(str(path), "export.zip")


def test_stepone_protocol_with_a_dtd_is_refused():
    with pytest.raises(ValueError, match="document type declaration"):
        plan_reads(_utf16(DOCTYPE_ENTITY.format(root="TCProtocol")))


def test_malformed_xml_is_a_value_error():
    with pytest.raises(ValueError, match="not well-formed"):
        safe_xml.parse_xml_bytes(b"<a><b></a>", "The document")


# --- plate and read limits ----------------------------------------------


def test_plate_declaring_one_thousand_square_is_refused(tmp_path):
    eds = build_quantstudio_eds(QuantStudioOptions(rows=1000, cols=1000))
    path = _write(tmp_path, "huge.eds", eds)
    with pytest.raises(ValueError, match="up to 384 wells"):
        detect_and_parse(path, "huge.eds")


@pytest.mark.parametrize("dims", [(8, 12), (16, 24)])
def test_supported_plate_formats_are_accepted(dims):
    xml = f"<PlateTypeID>TYPE_{dims[0]}X{dims[1]}</PlateTypeID>".encode()
    assert _parse_plate_dims(xml) == dims


@pytest.mark.parametrize("dims", [(32, 48), (1, 385), (17, 24), (0, 12)])
def test_plates_beyond_the_largest_format_are_refused(dims):
    xml = f"<PlateTypeID>TYPE_{dims[0]}X{dims[1]}</PlateTypeID>".encode()
    with pytest.raises(ValueError, match="up to 384 wells"):
        _parse_plate_dims(xml)


def _multicomponent(flags: str, well: int = 0, values: str = "[1.0, 2.0]") -> bytes:
    return (
        f"<Root><TCStageFlags>{flags}</TCStageFlags>"
        f'<DyeData WellIndex="{well}"><DyeList>[FAM, VIC]</DyeList></DyeData>'
        f'<SignalData WellIndex="{well}"><CycleData>{values}</CycleData>'
        f"<CycleData>{values}</CycleData></SignalData></Root>"
    ).encode()


def test_multicomponent_read_counts_are_bounded():
    _parse_multicomponent(_multicomponent("[" + ",".join(["5"] * 1000) + "]"))
    with pytest.raises(ValueError, match="more than 1000 values"):
        _parse_multicomponent(_multicomponent("[" + ",".join(["5"] * 1001) + "]"))
    with pytest.raises(ValueError, match="more than 1000 values"):
        _parse_multicomponent(
            _multicomponent("[1]", values="[" + ",".join(["1.0"] * 1001) + "]")
        )


@pytest.mark.parametrize("well", [-1, 384, 10_000_000])
def test_multicomponent_well_index_beyond_the_plate_is_refused(well):
    with pytest.raises(ValueError, match="plates of up to 384 wells"):
        _parse_multicomponent(_multicomponent("[1]", well=well))


def test_pcrd_plate_beyond_the_largest_format_is_refused(tmp_path, monkeypatch):
    root = safe_xml.parse_xml_bytes(
        b'<experimentalData2><plateSetup2 rows="1000" columns="1000"/></experimentalData2>',
        "doc",
    )
    monkeypatch.setattr(pcrd_raw, "_extract_xml", lambda _p: root)
    with pytest.raises(ValueError, match="up to 384 wells"):
        pcrd_raw.parse_pcrd("unused.pcrd")


def test_pcrd_read_count_is_bounded():
    reads = "<plateRead/>" * 1001
    root = safe_xml.parse_xml_bytes(
        f"<r><runData><plateReadDataVector>{reads}</plateReadDataVector></runData></r>".encode(),
        "doc",
    )
    with pytest.raises(ValueError, match="at most 1000"):
        pcrd_raw._parse_plate_reads(root, {"FAM": 0}, {0})


def _amp_xml(rows: int, wells: int) -> str:
    cells = "".join(f"<A{i + 1}>1</A{i + 1}>" for i in range(wells))
    body = "".join(f"<Row><Cycle>{r + 1}</Cycle>{cells}</Row>" for r in range(rows))
    return f"<FAM>{body}</FAM>"


def test_cfx_amplification_dimensions_are_bounded(tmp_path):
    ok = tmp_path / "ok.xml"
    ok.write_text(_amp_xml(3, 3))
    dye, cycles, data = cfx_xml_parser._parse_amplification_xml(str(ok))
    assert (dye, cycles, sorted(data)) == ("FAM", [1, 2, 3], ["A1", "A2", "A3"])

    too_many_reads = tmp_path / "reads.xml"
    too_many_reads.write_text(_amp_xml(1001, 1))
    with pytest.raises(ValueError, match="at most 1000"):
        cfx_xml_parser._parse_amplification_xml(str(too_many_reads))

    too_many_wells = tmp_path / "wells.xml"
    too_many_wells.write_text(_amp_xml(1, 385))
    with pytest.raises(ValueError, match="plates of up to 384 wells"):
        cfx_xml_parser._parse_amplification_xml(str(too_many_wells))


# --- archive directory ---------------------------------------------------


def _zip_bytes(entries: int) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i in range(entries):
            zf.writestr(f"e{i}.xml", b"")
    return buf.getvalue()


def _patch_eocd(data: bytes, *, entries=None, directory_size=None) -> bytes:
    pos = data.rfind(b"PK\x05\x06")
    fields = list(struct.unpack("<4s4H2LH", data[pos : pos + 22]))
    if entries is not None:
        fields[3] = fields[4] = entries
    if directory_size is not None:
        fields[5] = directory_size
    return data[:pos] + struct.pack("<4s4H2LH", *fields) + data[pos + 22 :]


def _as_zip64(data: bytes, entries: int, directory_size: int) -> bytes:
    """Rewrite the end record so the real totals live in the ZIP64 records."""
    pos = data.rfind(b"PK\x05\x06")
    _, disk, cd_disk, _, _, _, offset, comment = struct.unpack(
        "<4s4H2LH", data[pos : pos + 22]
    )
    record = struct.pack(
        "<4sQ2H2L4Q", b"PK\x06\x06", 44, 45, 45, 0, 0, entries, entries,
        directory_size, offset,
    )
    locator = struct.pack("<4sLQL", b"PK\x06\x07", 0, pos + 0, 1)
    end = struct.pack(
        "<4s4H2LH", b"PK\x05\x06", 0, 0, 0xFFFF, 0xFFFF, 0xFFFFFFFF, 0xFFFFFFFF, 0
    )
    return data[:pos] + record + locator + end


@pytest.mark.parametrize("ext", [".eds", ".pcrd", ".zip", ".xlsx"])
def test_archive_with_too_many_entries_is_refused_before_it_is_opened(
    tmp_path, monkeypatch, ext
):
    path = _write(tmp_path, f"many{ext}", _zip_bytes(600))

    def refuse(*_a, **_k):
        raise AssertionError("the archive directory was parsed")

    monkeypatch.setattr(zipfile, "ZipFile", refuse)
    with pytest.raises(ValueError, match=r"too many entries \(600 > 500\)"):
        detect_and_parse(path, f"many{ext}")


def test_zip64_end_record_is_read(tmp_path):
    data = _as_zip64(_zip_bytes(3), entries=600, directory_size=600 * 50)
    path = _write(tmp_path, "z64.zip", data)
    with pytest.raises(ValueError, match="too many entries"):
        safe_xml.check_zip_directory(path)


def test_zip64_end_record_within_limits_is_accepted(tmp_path):
    data = _as_zip64(_zip_bytes(3), entries=3, directory_size=3 * 50)
    safe_xml.check_zip_directory(_write(tmp_path, "z64.zip", data))


def test_oversize_central_directory_is_refused(tmp_path):
    data = _patch_eocd(_zip_bytes(3), directory_size=5 * 1024 * 1024)
    with pytest.raises(ValueError, match="directory is larger"):
        safe_xml.check_zip_directory(_write(tmp_path, "dir.zip", data))


def test_ordinary_archives_and_non_archives_pass_the_directory_check(tmp_path):
    safe_xml.check_zip_directory(_write(tmp_path, "ok.zip", _zip_bytes(500)))
    safe_xml.check_zip_directory(_write(tmp_path, "text.zip", b"not an archive"))
    safe_xml.check_zip_directory(_write(tmp_path, "empty.zip", b""))


def test_archive_with_a_comment_is_measured_from_its_end_record(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i in range(600):
            zf.writestr(f"e{i}.xml", b"")
        zf.comment = b"c" * 4000
    with pytest.raises(ValueError, match="too many entries"):
        safe_xml.check_zip_directory(_write(tmp_path, "c.zip", buf.getvalue()))


# --- CFX workbook --------------------------------------------------------


def _cfx_workbook(path: str, rows: int = 2) -> str:
    wb = openpyxl.Workbook()
    fam = wb.active
    fam.title = "FAM"
    hex_ = wb.create_sheet("HEX")
    for ws in (fam, hex_):
        ws.append([None, "Cycle", "A1"])
        for cycle in range(1, rows):
            ws.append([None, cycle, 1000.0 + cycle])
    info = wb.create_sheet("Run Information")
    info.append(["File Name", "run.pcrd"])
    wb.save(path)
    return path


def test_cfx_workbook_still_parses(tmp_path):
    path = _cfx_workbook(str(tmp_path / "cfx.xlsx"), rows=4)
    unified = detect_and_parse(path, "cfx.xlsx")
    assert unified.wells == ["A1"]
    assert unified.cycles == [1, 2, 3]


def test_cfx_detection_loads_the_workbook_once_and_read_only(tmp_path, monkeypatch):
    path = _cfx_workbook(str(tmp_path / "cfx.xlsx"))
    calls: list[dict] = []
    real = openpyxl.load_workbook

    def counting(*args, **kwargs):
        calls.append(kwargs)
        return real(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", counting)
    monkeypatch.setattr(cfx_opus, "parse_cfx_opus", lambda _p: "parsed")
    assert detector._handle_cfx_opus(path, "cfx.xlsx") == "parsed"
    assert len(calls) == 1
    assert calls[0].get("read_only") is True


def test_cfx_sheet_beyond_the_row_limit_is_refused(tmp_path, monkeypatch):
    path = _cfx_workbook(str(tmp_path / "cfx.xlsx"), rows=20)
    monkeypatch.setattr(detector, "MAX_XLSX_SHEET_ROWS", 10)
    with pytest.raises(ValueError, match="larger than an instrument export can be"):
        detect_and_parse(path, "cfx.xlsx")


def test_cfx_sheet_beyond_the_column_limit_is_refused(tmp_path, monkeypatch):
    path = _cfx_workbook(str(tmp_path / "cfx.xlsx"))
    monkeypatch.setattr(detector, "MAX_XLSX_SHEET_COLS", 2)
    with pytest.raises(ValueError, match="larger than an instrument export can be"):
        detect_and_parse(path, "cfx.xlsx")


def test_unreadable_workbook_keeps_its_message(tmp_path):
    path = tmp_path / "broken.xlsx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("hello.txt", "x")
    with pytest.raises(ValueError, match="Could not read this .xlsx file"):
        detect_and_parse(str(path), "broken.xlsx")
